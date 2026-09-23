"""Elaborazione: i valori ufficiali Istat per Italia, Nord, Centro e Mezzogiorno.

Il sito carica dall'appendice statistica del BES solo le 20 regioni
(`scripts/update_bes_regions.py` scarta il resto). Ma lo stesso file porta
anche le righe Italia, Nord, Centro e Mezzogiorno, calcolate dall'Istat sui
dati di base: sono medie **pesate**, e sono le uniche che un articolo puo'
chiamare "media italiana" o "valore del Mezzogiorno". La media semplice delle
regioni tratta il Molise come la Lombardia, e su alcuni indicatori sbaglia di
molto (i posti letto del Centro sono 60 per 10.000 abitanti, ma nella media
semplice "Centro-Nord" sparivano dentro le regioni del Nord).

Metodo: nessuna rielaborazione. Si leggono le righe con SESSO = Totale e
TERRITORIO in Italia, Nord, Centro, Mezzogiorno del file
`indicatori_regione_sesso.xlsx` dentro l'appendice statistica.

    bin/py -m scripts.trend_articles.derive_bes_areas SDG-311 01SAL005 ...

Scrive `data/derived/bes_areas_<codice>.csv` e `.json`, che `dossier.py`
legge come indicatore `ext:bes_areas_<codice>`.
"""

from __future__ import annotations

import argparse
import io
import sys
import zipfile

import requests
from openpyxl import load_workbook

from scripts.trend_articles import common

APPENDIX = "https://www.istat.it/wp-content/uploads/2026/05/APPENDICE-STATISTICA-2.zip"


def read_areas(codes: set[str]) -> dict[str, dict]:
    response = requests.get(APPENDIX, timeout=180)
    response.raise_for_status()
    with zipfile.ZipFile(io.BytesIO(response.content)) as archive:
        name = next(n for n in archive.namelist() if n.endswith("indicatori_regione_sesso.xlsx"))
        workbook = load_workbook(io.BytesIO(archive.read(name)), read_only=True)
    rows = workbook[workbook.sheetnames[0]].iter_rows(values_only=True)
    header = next(rows)
    year_columns = [(i, c) for i, c in enumerate(header) if isinstance(c, int)]
    result: dict[str, dict] = {}
    for r in rows:
        if r[1] not in codes or r[3] != "Totale" or r[4] not in common.AREAS:
            continue
        entry = result.setdefault(r[1], {"name": r[2], "unit": r[5], "source": r[6], "values": {}})
        for i, year in year_columns:
            v = common.parse_number(str(r[i])) if r[i] not in (None, "") else None
            if v is not None:
                entry["values"].setdefault(r[4], {})[year] = v
    missing = codes - set(result)
    if missing:
        raise KeyError(f"codici assenti dall'appendice BES ({APPENDIX}): {sorted(missing)}")
    return result


def main(argv=None) -> int:
    parser = argparse.ArgumentParser(description=__doc__.split("\n")[0])
    parser.add_argument("codes", nargs="+")
    args = parser.parse_args(argv)
    common.DERIVED_DIR.mkdir(parents=True, exist_ok=True)
    for code, entry in read_areas(set(args.codes)).items():
        base = common.DERIVED_DIR / f"bes_areas_{code}"
        with base.with_suffix(".csv").open("w", encoding="utf-8") as file:
            file.write("territory,year,value\n")
            for area in common.AREAS:
                for year, v in sorted(entry["values"].get(area, {}).items()):
                    file.write(f"{area},{year},{v}\n")
        common.write_json(base.with_suffix(".json"), {
            "name": f"{entry['name']}, Italia e ripartizioni",
            "unit": entry["unit"],
            "source": "Istat",
            "archive": "Benessere equo e sostenibile, aggiornamento intermedio 2026, valori ufficiali per ripartizione",
            "source_url": APPENDIX,
            "license": "CC BY 4.0",
            "method": "Valori Istat per Italia, Nord, Centro e Mezzogiorno, senza rielaborazione (medie calcolate dall'Istat sui dati di base, non medie semplici delle regioni).",
            "script": "scripts/trend_articles/derive_bes_areas.py",
            "bes_code": code,
        })
        print(code, {a: (min(v), max(v)) for a, v in entry["values"].items()})
    return 0


if __name__ == "__main__":
    sys.exit(main())
