#!/usr/bin/env python3
"""Estrazione dati ACI (Automobile Club d'Italia, Autoritratto 2025).

Estrae due indicatori provinciali e regionali per l'anno 2025:
1. ACI_AUTO_ANTE_2009: quota autovetture immatricolate fino al 2009 (%)
2. ACI_AUTO_ALIMENTAZIONE_ALTERNATIVA: quota autovetture non benzina/gasolio (%)
"""

import argparse
import csv
import os
import re
import sys
import unicodedata
import urllib.request
import zipfile
import openpyxl

URL_ACI_2025 = "https://aci.gov.it/app/uploads/2026/06/Autoritratto2025_Parco_veicolare.zip"
CACHE_DIR = "lavoro/cache"
ZIP_CACHE_PATH = os.path.join(CACHE_DIR, "Autoritratto2025_Parco_veicolare.zip")
XLSX_CACHE_PATH = os.path.join(CACHE_DIR, "Parco_veicolare_2025.xlsx")

OUT_DATA_PATH = "app/static/data/nuovi/aci.csv"
OUT_MANIFEST_PATH = "app/static/data/nuovi/aci_manifest.csv"
PROVINCE_CODES_PATH = "app/static/data/province_codes.csv"

REGION_MAP = {
    "PIEMONTE": "piemonte",
    "VALLE D'AOSTA": "valle-d-aosta",
    "LOMBARDIA": "lombardia",
    "TRENTINO ALTO ADIGE": "trentino-alto-adige",
    "TRENTINO-ALTO ADIGE": "trentino-alto-adige",
    "VENETO": "veneto",
    "FRIULI-VENEZIA GIULIA": "friuli-venezia-giulia",
    "FRIULI VENEZIA GIULIA": "friuli-venezia-giulia",
    "LIGURIA": "liguria",
    "EMILIA-ROMAGNA": "emilia-romagna",
    "EMILIA ROMAGNA": "emilia-romagna",
    "TOSCANA": "toscana",
    "UMBRIA": "umbria",
    "MARCHE": "marche",
    "LAZIO": "lazio",
    "ABRUZZO": "abruzzo",
    "MOLISE": "molise",
    "CAMPANIA": "campania",
    "PUGLIA": "puglia",
    "BASILICATA": "basilicata",
    "CALABRIA": "calabria",
    "SICILIA": "sicilia",
    "SARDEGNA": "sardegna",
}

PROVINCE_ALIASES = {
    "BARLETTA TRANI": "barletta-andria-trani",
    "BARLETTA-ANDRIA-TRANI": "barletta-andria-trani",
    "BARLETTA ANDRIA TRANI": "barletta-andria-trani",
    "MONZA BRIANZA": "monza-e-della-brianza",
    "MONZA E DELLA BRIANZA": "monza-e-della-brianza",
    "REGGIO CALABRIA": "reggio-calabria",
    "REGGIO DI CALABRIA": "reggio-calabria",
    "REGGIO EMILIA": "reggio-emilia",
    "REGGIO NELL'EMILIA": "reggio-emilia",
    "REGGIO NELL’EMILIA": "reggio-emilia",
    "FORLI' CESENA": "forli-cesena",
    "FORLI'-CESENA": "forli-cesena",
    "FORLI CESENA": "forli-cesena",
    "MASSA CARRARA": "massa-carrara",
    "MASSA-CARRARA": "massa-carrara",
    "PESARO URBINO": "pesaro-e-urbino",
    "PESARO E URBINO": "pesaro-e-urbino",
    "VERBANO CUSIO OSSOLA": "verbano-cusio-ossola",
    "VERBANO-CUSIO-OSSOLA": "verbano-cusio-ossola",
    "VALLE D'AOSTA": "aosta",
    "AOSTA": "aosta",
    "BOLZANO": "bolzano",
    "BOLZANO-BOZEN": "bolzano",
    "BOLZANO/BOZEN": "bolzano",
    "L'AQUILA": "l-aquila",
    "L’AQUILA": "l-aquila",
    "AQUILA": "l-aquila",
    "SUD SARDEGNA": "sud-sardegna",
}


def slugify(val):
    val = unicodedata.normalize("NFKD", val or "").encode("ascii", "ignore").decode("ascii")
    val = val.lower().replace("'", " ")
    return re.sub(r"[^a-z0-9]+", "-", val).strip("-")


def load_province_keys(path):
    prov_keys = {}
    with open(path, "r", encoding="utf-8") as f:
        reader = csv.DictReader(f, delimiter=";")
        for row in reader:
            prov_keys[slugify(row["name"])] = row["province_key"]
    return prov_keys


def ensure_file(offline=False):
    os.makedirs(CACHE_DIR, exist_ok=True)
    if os.path.exists(XLSX_CACHE_PATH):
        return
    if os.path.exists(ZIP_CACHE_PATH):
        with zipfile.ZipFile(ZIP_CACHE_PATH, "r") as z:
            z.extract("Parco_veicolare_2025.xlsx", CACHE_DIR)
        return
    if offline:
        raise FileNotFoundError(f"File non presente in cache: {XLSX_CACHE_PATH}")
    print(f"Scaricamento file ACI da {URL_ACI_2025}...")
    req = urllib.request.Request(URL_ACI_2025, headers={"User-Agent": "Mozilla/5.0"})
    with urllib.request.urlopen(req, timeout=30) as resp:
        data = resp.read()
        with open(ZIP_CACHE_PATH, "wb") as f:
            f.write(data)
    with zipfile.ZipFile(ZIP_CACHE_PATH, "r") as z:
        z.extract("Parco_veicolare_2025.xlsx", CACHE_DIR)


def extract_aci_data(xlsx_path, prov_map):
    wb = openpyxl.load_workbook(xlsx_path, data_only=True)
    
    # 1. Indicatore ACI_AUTO_ANTE_2009 da foglio "28 AV Provincia anno"
    ws28 = wb["28 AV Provincia anno"]
    prov_ante = {}
    prov_tot28 = {}
    reg_ante = {}
    reg_tot28 = {}
    
    current_reg = None
    
    for r in range(4, ws28.max_row + 1):
        reg_val = ws28.cell(r, 1).value
        prov_val = ws28.cell(r, 2).value
        
        if reg_val:
            current_reg = str(reg_val).strip()
        
        if not prov_val:
            continue
        
        prov_name = str(prov_val).strip()
        if prov_name in ["Area Geografica", "ITALIA NORD-OCCIDENTALE", "ITALIA NORD-ORIENTALE",
                         "ITALIA CENTRALE", "ITALIA MERIDIONALE", "ITALIA INSULARE", "NON DEFINITO"]:
            continue
            
        fino_2009 = ws28.cell(r, 3).value or 0
        totale = ws28.cell(r, 12).value or 0
        
        if prov_name == "Totale":
            if current_reg and current_reg in REGION_MAP:
                rkey = REGION_MAP[current_reg]
                reg_ante[rkey] = reg_ante.get(rkey, 0) + fino_2009
                reg_tot28[rkey] = reg_tot28.get(rkey, 0) + totale
        else:
            pkey = PROVINCE_ALIASES.get(prov_name) or prov_map.get(slugify(prov_name))
            if pkey:
                prov_ante[pkey] = fino_2009
                prov_tot28[pkey] = totale

    # 2. Indicatore ACI_AUTO_ALIMENTAZIONE_ALTERNATIVA da fogli "36" e "27"
    ws27 = wb["27 AV Provincia cilindrata"]
    ws36 = wb["36 AVAltre Provincia cilindrata"]
    
    prov_tot27 = {}
    reg_tot27 = {}
    current_reg = None
    
    for r in range(4, ws27.max_row + 1):
        reg_val = ws27.cell(r, 1).value
        prov_val = ws27.cell(r, 2).value
        if reg_val:
            current_reg = str(reg_val).strip()
        if not prov_val:
            continue
        prov_name = str(prov_val).strip()
        if prov_name in ["Area Geografica", "ITALIA NORD-OCCIDENTALE", "ITALIA NORD-ORIENTALE",
                         "ITALIA CENTRALE", "ITALIA MERIDIONALE", "ITALIA INSULARE", "NON DEFINITO"]:
            continue
            
        totale = ws27.cell(r, 12).value or 0
        if prov_name == "Totale":
            if current_reg and current_reg in REGION_MAP:
                rkey = REGION_MAP[current_reg]
                reg_tot27[rkey] = reg_tot27.get(rkey, 0) + totale
        else:
            pkey = PROVINCE_ALIASES.get(prov_name) or prov_map.get(slugify(prov_name))
            if pkey:
                prov_tot27[pkey] = totale

    prov_altre = {}
    reg_altre = {}
    current_reg = None
    
    for r in range(4, ws36.max_row + 1):
        reg_val = ws36.cell(r, 1).value
        prov_val = ws36.cell(r, 2).value
        if reg_val:
            current_reg = str(reg_val).strip()
        if not prov_val:
            continue
        prov_name = str(prov_val).strip()
        if prov_name in ["Area Geografica", "ITALIA NORD-OCCIDENTALE", "ITALIA NORD-ORIENTALE",
                         "ITALIA CENTRALE", "ITALIA MERIDIONALE", "ITALIA INSULARE", "NON DEFINITO"]:
            continue
            
        altre = ws36.cell(r, 12).value or 0
        if prov_name == "Totale":
            if current_reg and current_reg in REGION_MAP:
                rkey = REGION_MAP[current_reg]
                reg_altre[rkey] = reg_altre.get(rkey, 0) + altre
        else:
            pkey = PROVINCE_ALIASES.get(prov_name) or prov_map.get(slugify(prov_name))
            if pkey:
                prov_altre[pkey] = altre

    # Costruzione righe CSV
    rows = []
    
    # Ante 2009 Province
    for pkey, num in sorted(prov_ante.items()):
        den = prov_tot28.get(pkey, 0)
        if den > 0:
            val = round((num / den) * 100, 1)
            rows.append({
                "indicator_id": "ACI_AUTO_ANTE_2009",
                "level": "provincia",
                "territory_key": pkey,
                "year": 2025,
                "value": f"{val:.1f}"
            })
            
    # Ante 2009 Regioni
    for rkey, num in sorted(reg_ante.items()):
        den = reg_tot28.get(rkey, 0)
        if den > 0:
            val = round((num / den) * 100, 1)
            rows.append({
                "indicator_id": "ACI_AUTO_ANTE_2009",
                "level": "regione",
                "territory_key": rkey,
                "year": 2025,
                "value": f"{val:.1f}"
            })

    # Alimentazione alternativa Province
    for pkey, num in sorted(prov_altre.items()):
        den = prov_tot27.get(pkey, 0)
        if den > 0:
            val = round((num / den) * 100, 1)
            rows.append({
                "indicator_id": "ACI_AUTO_ALIMENTAZIONE_ALTERNATIVA",
                "level": "provincia",
                "territory_key": pkey,
                "year": 2025,
                "value": f"{val:.1f}"
            })

    # Alimentazione alternativa Regioni
    for rkey, num in sorted(reg_altre.items()):
        den = reg_tot27.get(rkey, 0)
        if den > 0:
            val = round((num / den) * 100, 1)
            rows.append({
                "indicator_id": "ACI_AUTO_ALIMENTAZIONE_ALTERNATIVA",
                "level": "regione",
                "territory_key": rkey,
                "year": 2025,
                "value": f"{val:.1f}"
            })

    return rows


def write_outputs(rows):
    os.makedirs(os.path.dirname(OUT_DATA_PATH), exist_ok=True)
    with open(OUT_DATA_PATH, "w", encoding="utf-8", newline="") as f:
        writer = csv.DictWriter(f, fieldnames=["indicator_id", "level", "territory_key", "year", "value"], delimiter=";")
        writer.writeheader()
        writer.writerows(rows)
        
    manifest_rows = [
        {
            "indicator_id": "ACI_AUTO_ANTE_2009",
            "name": "Quota di autovetture immatricolate fino al 2009",
            "institution": "Automobile Club d'Italia",
            "unit": "%",
            "decimals": 1,
            "direction": "lower_better",
            "theme": "Trasporti e mobilità",
            "source_url": "https://aci.gov.it/attivita-e-progetti/studi-e-ricerche/autoritratto/",
            "method_url": "https://aci.gov.it/attivita-e-progetti/studi-e-ricerche/autoritratto/",
            "license": "CC BY 4.0",
            "license_quote": "liberamente fruibili da chiunque nel rispetto dei termini previsti dalla licenza di utilizzo Creative Commons CC-BY 4.0",
            "year_min": 2025,
            "year_max": 2025,
            "n_provincia": 107,
            "n_regione": 20,
            "note": "Quota percentuale di autovetture del parco veicolare al 31/12 immatricolate fino al 2009. Elaborazione su dati ACI Autoritratto 2025."
        },
        {
            "indicator_id": "ACI_AUTO_ALIMENTAZIONE_ALTERNATIVA",
            "name": "Quota di autovetture con alimentazione alternativa a benzina e gasolio",
            "institution": "Automobile Club d'Italia",
            "unit": "%",
            "decimals": 1,
            "direction": "higher_better",
            "theme": "Trasporti e mobilità",
            "source_url": "https://aci.gov.it/attivita-e-progetti/studi-e-ricerche/autoritratto/",
            "method_url": "https://aci.gov.it/attivita-e-progetti/studi-e-ricerche/autoritratto/",
            "license": "CC BY 4.0",
            "license_quote": "liberamente fruibili da chiunque nel rispetto dei termini previsti dalla licenza di utilizzo Creative Commons CC-BY 4.0",
            "year_min": 2025,
            "year_max": 2025,
            "n_provincia": 107,
            "n_regione": 20,
            "note": "Quota percentuale di autovetture con alimentazione diversa da benzina e gasolio (ibride, GPL, metano, elettriche). La quota alta di alcune province (es. Firenze) dipende anche dalla presenza di flotte aziendali e noleggi."
        }
    ]
    
    with open(OUT_MANIFEST_PATH, "w", encoding="utf-8", newline="") as f:
        fieldnames = ["indicator_id", "name", "institution", "unit", "decimals", "direction", "theme",
                      "source_url", "method_url", "license", "license_quote", "year_min", "year_max",
                      "n_provincia", "n_regione", "note"]
        writer = csv.DictWriter(f, fieldnames=fieldnames, delimiter=";")
        writer.writeheader()
        writer.writerows(manifest_rows)


def main():
    parser = argparse.ArgumentParser(description="Estrattore dati ACI")
    parser.add_argument("--offline", action="store_true", help="Usa solo i file in cache locale")
    args = parser.parse_args()

    ensure_file(offline=args.offline)
    prov_map = load_province_keys(PROVINCE_CODES_PATH)
    rows = extract_aci_data(XLSX_CACHE_PATH, prov_map)
    write_outputs(rows)
    print(f"Estratte {len(rows)} righe per ACI. Salvare in {OUT_DATA_PATH} e {OUT_MANIFEST_PATH}")


if __name__ == "__main__":
    main()
