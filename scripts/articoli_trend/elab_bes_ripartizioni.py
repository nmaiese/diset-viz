"""Elaborazione: i valori ufficiali Istat per Italia, Nord, Centro e Mezzogiorno.

Il sito carica dall'appendice statistica del BES solo le 20 regioni
(`scripts/update_bes_regions.py` scarta il resto). Ma lo stesso file porta
anche le righe Italia, Nord, Centro e Mezzogiorno, calcolate dall'Istat sui
dati di base: sono medie **pesate**, e sono le uniche che un articolo puo'
chiamare "media italiana" o "valore del Mezzogiorno". La media semplice delle
regioni tratta il Molise come la Lombardia, e su alcuni indicatori sbaglia di
molto (i posti letto del Centro sono 60 per 10.000 abitanti, ma nella media
semplice "Centro-Nord" sparivano dentro le regioni del Nord).

Metodo: nessuna rielaborazione. Si leggono le righe con SESSO = Totale e
TERRITORIO in Italia, Nord, Centro, Mezzogiorno del file
`indicatori_regione_sesso.xlsx` dentro l'appendice statistica.

    bin/py -m scripts.articoli_trend.elab_bes_ripartizioni SDG-311 01SAL005 ...

Scrive `data/elaborazioni/bes_ripartizioni_<codice>.csv` e `.json`, che
`dossier.py` legge come indicatore `ext:bes_ripartizioni_<codice>`.
"""

from __future__ import annotations

import argparse
import io
import sys
import zipfile

import requests
from openpyxl import load_workbook

from scripts.articoli_trend import comuni

APPENDICE = "https://www.istat.it/wp-content/uploads/2026/05/APPENDICE-STATISTICA-2.zip"
AREE = ("Italia", "Nord", "Centro", "Mezzogiorno")
CARTELLA = comuni.RADICE / "data" / "elaborazioni"


def leggi(codici: set[str]) -> dict[str, dict]:
    grezzo = requests.get(APPENDICE, timeout=180).content
    with zipfile.ZipFile(io.BytesIO(grezzo)) as z:
        nome = next(n for n in z.namelist() if n.endswith("indicatori_regione_sesso.xlsx"))
        wb = load_workbook(io.BytesIO(z.read(nome)), read_only=True)
    ws = wb[wb.sheetnames[0]]
    righe = ws.iter_rows(values_only=True)
    testa = next(righe)
    anni = [(i, c) for i, c in enumerate(testa) if isinstance(c, int)]
    esito: dict[str, dict] = {}
    for r in righe:
        if r[1] not in codici or r[3] != "Totale" or r[4] not in AREE:
            continue
        e = esito.setdefault(r[1], {"nome": r[2], "unita": r[5], "fonte": r[6], "valori": {}})
        for i, anno in anni:
            v = comuni.numero(str(r[i])) if r[i] not in (None, "") else None
            if v is not None:
                e["valori"].setdefault(r[4], {})[anno] = v
    return esito


def main(argv=None) -> int:
    parser = argparse.ArgumentParser(description=__doc__.split("\n")[0])
    parser.add_argument("codici", nargs="+")
    args = parser.parse_args(argv)
    CARTELLA.mkdir(parents=True, exist_ok=True)
    for codice, e in leggi(set(args.codici)).items():
        base = CARTELLA / f"bes_ripartizioni_{codice}"
        with base.with_suffix(".csv").open("w", encoding="utf-8") as f:
            f.write("territorio,anno,valore\n")
            for area in AREE:
                for anno, v in sorted(e["valori"].get(area, {}).items()):
                    f.write(f"{area},{anno},{v}\n")
        comuni.scrivi_json(base.with_suffix(".json"), {
            "nome": f"{e['nome']}, Italia e ripartizioni",
            "unita": e["unita"],
            "fonte": "Istat",
            "archivio": "Benessere equo e sostenibile, aggiornamento intermedio 2026, valori ufficiali per ripartizione",
            "source_url": APPENDICE,
            "licenza": "CC BY 4.0",
            "method": "Valori Istat per Italia, Nord, Centro e Mezzogiorno, senza rielaborazione (medie calcolate dall'Istat sui dati di base, non medie semplici delle regioni).",
            "script": "scripts/articoli_trend/elab_bes_ripartizioni.py",
            "codice_bes": codice,
        })
        print(codice, {a: (min(v), max(v)) for a, v in e["valori"].items()})
    return 0


if __name__ == "__main__":
    sys.exit(main())
